"""Downloads generated from actual outputs, never from evaluation references."""
import csv,io,json,zipfile
from docx import Document

def json_bytes(value): return json.dumps(value,indent=2,ensure_ascii=False).encode('utf-8')

def meeting_text(bundle):
    o=bundle['output'];r=bundle['record']
    lines=[o['meeting_title'],'Generated meeting notes — review required',
        f"Model: {r['requested_model']} | Prompt: {r['prompt_type']}",
        'SUMMARY',o['summary'],'KEY TOPICS',*o['key_topics'],'DECISIONS']
    for i in o['decisions']:
        lines += [i['decision'],*[f'Evidence: {q}' for q in i['evidence']]]
    lines.append('ACTIONS')
    for i in o['action_items']:
        lines += [i['action'],f"Owner: {i['owner'] or 'Not specified'}",f"Deadline: {i['deadline'] or 'Not specified'}",
            f"Condition: {i['condition'] or 'None specified'}",*[f'Evidence: {q}' for q in i['evidence']]]
    lines += ['UNRESOLVED ISSUES',*o['unresolved_issues'],'AMBIGUITIES',*o['ambiguities'],
        'Schema validity and quotation matches do not establish factual accuracy.']
    return '\n\n'.join(lines)

def email_draft(bundle):
    o=bundle['output']
    lines=[f"Subject: Follow-up — {o['meeting_title']}",'',
        'DRAFT — verify against the transcript before sending.','',
        'Hello everyone,','', 'Please check the following generated meeting notes.','',
        'Decisions:']
    lines += [f"- {i['decision']}" for i in o['decisions']]
    lines += ['', 'Actions:']
    for i in o['action_items']:
        line=f"- {i['action']} | Owner: {i['owner'] or 'To confirm'} | Deadline: {i['deadline'] or 'Not specified'}"
        if i['condition']: line+=f" | Condition: {i['condition']}"
        lines.append(line)
    lines += ['','Please confirm any missing ownership, dates or corrections.','','Kind regards,','[Sender]']
    return '\n'.join(lines)

def docx_bytes(bundle):
    doc=Document();doc.add_heading('AI Meeting Intelligence',0)
    for line in meeting_text(bundle).split('\n\n'): doc.add_paragraph(line)
    out=io.BytesIO();doc.save(out);return out.getvalue()

def metrics_csv(report):
    out=io.StringIO();writer=csv.DictWriter(out,fieldnames=['metric','numerator','denominator','percentage'])
    writer.writeheader();writer.writerows(report['metrics']);return out.getvalue().encode('utf-8')

def evidence_zip(project):
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('meeting_metadata.json',json_bytes({k:v for k,v in project.items() if k not in {'raw','clean','reviewed','attempts','evaluations'}}))
        for key in ('raw','clean','reviewed'):
            if key in project: z.writestr(f'transcripts/{key}.txt',project[key])
        for b in project.get('attempts',[]):
            name=b['record']['attempt_id'];z.writestr(f'attempts/{name}/evidence.json',json_bytes(b))
            if b['output']:
                z.writestr(f'attempts/{name}/notes.txt',meeting_text(b))
                z.writestr(f'attempts/{name}/notes.docx',docx_bytes(b))
                z.writestr(f'attempts/{name}/email_draft.txt',email_draft(b))
        for key,r in project.get('evaluations',{}).items():
            z.writestr(f'evaluation/{key}.json',json_bytes(r));z.writestr(f'evaluation/{key}.csv',metrics_csv(r))
    return out.getvalue()
